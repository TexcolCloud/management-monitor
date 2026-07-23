from __future__ import annotations

import argparse
import json
import logging
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from common.dataset_io import load_rows
from common.logging_utils import setup_logging
from common.runtime_config import load_runtime_config
from common.spreadsheet_safety import is_spreadsheet_formula
from common.workflow_paths import (
    DEFAULT_CAPTURE_ROOT,
    project_path,
    resolve_capture_data_dir,
    safe_slug,
)
LOGGER = logging.getLogger("workorder_daily_manage")
DEFAULT_EXCEL_NAME = "daily-management-table-no-operation.xlsx"
DEFAULT_EXPORT_ROOT = load_runtime_config().export_root
HEADERS = ["单据编号", "单位名称", "类型", "主题", "创建人", "创建时间"]
FIELD_MAP = {
    "单据编号": "safetyCode",
    "单位名称": "companyName",
    "类型": "safetyType",
    "主题": "theme",
    "创建人": "createBy",
    "创建时间": "createTime",
}
INVALID_XML_CHARS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
INVALID_SHEET_NAME_CHARS = re.compile(r"[\[\]:*?/\\]")
EXCEL_CELL_LIMIT = 32767


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export captured daily-management rows to an Excel table."
    )
    parser.add_argument("--data-dir", type=Path, default=None)
    parser.add_argument("--capture-root", type=Path, default=DEFAULT_CAPTURE_ROOT)
    parser.add_argument("--input", type=Path, default=None)
    parser.add_argument("--out", type=Path, default=None)
    parser.add_argument("--export-root", type=Path, default=DEFAULT_EXPORT_ROOT)
    parser.add_argument("--sheet-name", default="日常管理")
    return parser.parse_args()


def resolve_paths(args: argparse.Namespace) -> tuple[Path, Path, Path]:
    if args.input:
        input_path = project_path(args.input).resolve()
        data_dir = input_path.parent
    else:
        data_dir = resolve_capture_data_dir(args.data_dir, args.capture_root)
        input_path = data_dir / "management-api-all.json"
    output_path = (
        project_path(args.out).resolve()
        if args.out
        else project_path(args.export_root).resolve() / safe_slug(data_dir.name) / DEFAULT_EXCEL_NAME
    )
    return data_dir, input_path, output_path


def clean_text(value) -> str:
    if value is None:
        return ""
    text = json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)
    text = INVALID_XML_CHARS.sub("", text).replace("\r\n", "\n").replace("\r", "\n")
    return text[:EXCEL_CELL_LIMIT]


def validate_sheet_name(value: str) -> str:
    name = clean_text(value).strip()
    if not name:
        raise ValueError("Excel sheet name must not be empty")
    if len(name) > 31:
        raise ValueError("Excel sheet name must not exceed 31 characters")
    if INVALID_SHEET_NAME_CHARS.search(name) or name.startswith("'") or name.endswith("'"):
        raise ValueError(f"Invalid Excel sheet name: {value!r}")
    return name


def table_rows(rows: list[dict]) -> list[list[str]]:
    exported: list[list[str]] = []
    for row in rows:
        exported.append(
            [
                clean_text(row.get(FIELD_MAP["单据编号"])),
                clean_text(row.get(FIELD_MAP["单位名称"])),
                clean_text(row.get(FIELD_MAP["类型"])),
                clean_text(row.get(FIELD_MAP["主题"])),
                clean_text(row.get(FIELD_MAP["创建人"])),
                clean_text(row.get(FIELD_MAP["创建时间"])),
            ]
        )
    return exported


def style_worksheet(sheet: Worksheet) -> None:
    thin = Side(style="thin", color="D9DDE3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="EDEFF3")
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="2F4969")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=11)
            cell.alignment = Alignment(vertical="center")
            cell.border = border
    widths = [24, 38, 14, 38, 18, 22]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, sheet.max_row)}"


def write_xlsx(output_path: Path, sheet_name: str, rows: list[list[str]]) -> None:
    workbook = Workbook()
    workbook.properties.creator = "management-monitorDailyManage"
    sheet = workbook.active
    sheet.title = validate_sheet_name(sheet_name)
    sheet.append(HEADERS)
    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            text = clean_text(value)
            cell = sheet.cell(row=row_index, column=column_index, value=text)
            if is_spreadsheet_formula(text):
                cell.data_type = "s"
    style_worksheet(sheet)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    setup_logging()
    args = parse_args()
    try:
        data_dir, input_path, output_path = resolve_paths(args)
        exported_rows = table_rows(load_rows(input_path))
        write_xlsx(output_path, args.sheet_name, exported_rows)
    except (FileNotFoundError, ValueError) as exc:
        raise SystemExit(f"Export failed: {exc}") from exc
    LOGGER.info("Excel 导出完成: rows=%s file=%s", len(exported_rows), output_path.resolve())


if __name__ == "__main__":
    main()
