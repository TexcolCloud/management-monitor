from __future__ import annotations

import csv
import json
import tempfile
import unittest
from argparse import Namespace
from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock, Mock, patch

from openpyxl import load_workbook

from common.dataset_io import load_rows
from database.migrations import load_migrations
from common.runtime_config import load_runtime_config
from data_acquisition.capture_components import CaptureResult, CaptureStatistics, OutputWriter
from data_acquisition.capture_daily_range import run_capture
from database.postgres_store import (
    data_dir_records_complete,
    database_has_records,
    load_detail_batch,
    parse_create_time,
    prepare_records,
    rows_by_create_time,
)
from data_export.daily_management_excel import table_rows, validate_sheet_name, write_xlsx


class DatabaseAndExportTest(unittest.TestCase):
    def test_database_has_records_checks_target_table_after_migration(self) -> None:
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = (True,)
        config = {"enabled": True, "schema": "public", "table": "daily_manage"}
        with patch("database.postgres_store.load_config", return_value=config), patch(
            "database.postgres_store.connect", return_value=connection
        ), patch("database.postgres_store.ensure_table") as ensure:
            self.assertTrue(database_has_records())
        ensure.assert_called_once_with(connection, config)
        self.assertIn("SELECT EXISTS", cursor.execute.call_args.args[0])
        connection.close.assert_called_once()

    def test_rows_by_create_time_loads_incomplete_raw_rows_inclusive(self) -> None:
        connection = MagicMock()
        cursor = connection.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [
            ({"safetyCode": "CODE-1"},),
            ({"safetyCode": "CODE-2", "_captureError": "detail timeout"},),
        ]
        config = {"enabled": True, "schema": "public", "table": "daily_manage"}
        start = datetime(2026, 7, 1)
        end = datetime(2026, 7, 1, 23, 59, 59, 999999)
        with patch("database.postgres_store.load_config", return_value=config), patch(
            "database.postgres_store.connect", return_value=connection
        ), patch("database.postgres_store.ensure_table") as ensure:
            rows = rows_by_create_time(start, end)
        self.assertEqual(
            rows,
            [
                {"safetyCode": "CODE-1"},
                {"safetyCode": "CODE-2", "_captureError": "detail timeout"},
            ],
        )
        ensure.assert_called_once_with(connection, config)
        query, params = cursor.execute.call_args.args
        self.assertIn("create_time >= %s", query)
        self.assertIn("create_time <= %s", query)
        self.assertNotIn("_captureError", query)
        self.assertEqual(params, (start, end))
        connection.close.assert_called_once()

    def test_data_dir_records_complete_requires_each_captured_safety_code(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir)
            complete_row = {
                "id": "1",
                "safetyCode": "CODE-1",
                "createTime": "2026-07-23 10:00:00",
            }
            (data_dir / "management-api-all.json").write_text(
                json.dumps(
                    {
                        "statistics": {
                            "total_consistent": True,
                            "incomplete_reasons": [],
                        },
                        "rows": [complete_row],
                    }
                ),
                encoding="utf-8",
            )
            (data_dir / "management-api-details.json").write_text(
                json.dumps({"rows": [complete_row]}),
                encoding="utf-8",
            )
            (data_dir / "management-api-detail-failures.json").write_text(
                json.dumps({"rows": []}),
                encoding="utf-8",
            )
            (data_dir / "daily-attachments-manifest.json").write_text(
                json.dumps({"rows": []}),
                encoding="utf-8",
            )
            connection = MagicMock()
            config = {"enabled": True, "schema": "public", "table": "daily_manage"}
            with patch("database.postgres_store.load_config", return_value=config), patch(
                "database.postgres_store.connect", return_value=connection
            ), patch("database.postgres_store.ensure_table"), patch(
                "database.postgres_store.existing_safety_codes", return_value={"CODE-1"}
            ):
                self.assertTrue(data_dir_records_complete(data_dir))
            with patch("database.postgres_store.load_config", return_value=config), patch(
                "database.postgres_store.connect", return_value=connection
            ), patch("database.postgres_store.ensure_table"), patch(
                "database.postgres_store.existing_safety_codes", return_value=set()
            ):
                self.assertFalse(data_dir_records_complete(data_dir))
            (data_dir / "management-api-details.json").write_text(
                json.dumps({"rows": [{"theme": "missing code"}]}),
                encoding="utf-8",
            )
            self.assertFalse(data_dir_records_complete(data_dir))

    def test_data_dir_records_complete_rejects_partial_capture(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir)
            (data_dir / "management-api-all.json").write_text(
                json.dumps(
                    {
                        "statistics": {
                            "source_total_reported": 3,
                            "source_rows_fetched": 0,
                            "total_consistent": False,
                        }
                    }
                ),
                encoding="utf-8",
            )
            (data_dir / "management-api-details.json").write_text(
                json.dumps({"rows": []}),
                encoding="utf-8",
            )

            self.assertFalse(data_dir_records_complete(data_dir))

    def test_database_migrations_are_numbered_and_renderable(self) -> None:
        migrations = load_migrations()
        self.assertEqual([item.version for item in migrations], [1, 2, 3, 4, 5, 6])
        context = {
            "schema": '"public"',
            "table": '"public"."test"',
            "unique_constraint": '"uk_test_safety_code"',
        }
        initial = migrations[0].sql.format_map(context)
        outbox = migrations[1].sql.format_map(context)
        progress = migrations[2].sql.format_map(context)
        index = migrations[3].sql.format_map({**context, "create_time_index": '"idx_test_create_time"'})
        recovery = migrations[4].sql.format_map(context)
        reconciliation = migrations[5].sql.format_map(
            {**context, "create_time_index": '"idx_test_create_time"'}
        )
        self.assertIn("CREATE TABLE IF NOT EXISTS", initial)
        self.assertNotIn("CREATE UNIQUE INDEX", initial)
        self.assertIn("COMMENT ON TABLE", initial)
        self.assertIn("COMMENT ON COLUMN", initial)
        self.assertIn("workorder_feishu_notification_outbox", outbox)
        self.assertIn("card_sent_at", progress)
        self.assertIn("sent_image_paths", progress)
        self.assertIn("CREATE INDEX", index)
        self.assertIn("workorder_monitor_checkpoints", recovery)
        self.assertIn("next_attempt_at", recovery)
        self.assertIn("lease_owner", recovery)
        self.assertIn("ADD COLUMN IF NOT EXISTS", reconciliation)
        self.assertIn("duplicate safety_code", reconciliation)
        self.assertIn("ADD CONSTRAINT", reconciliation)

    def test_dataset_loader_rejects_non_object_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "rows.json"
            input_path.write_text(json.dumps({"rows": [{"id": 1}, "invalid"]}), encoding="utf-8")
            with self.assertRaises(ValueError):
                load_rows(input_path)

    def test_excel_company_name_uses_classifier_normalization(self) -> None:
        exported = table_rows([{"companyName": "示例分公司｜城区分公司"}])
        self.assertEqual(exported[0][1], "城区分公司")

    def test_parse_create_time_supports_timezone(self) -> None:
        source = "2026-07-21T10:00:00+08:00"
        expected = datetime.fromisoformat(source).astimezone().replace(tzinfo=None)
        self.assertEqual(parse_create_time(source), expected)

    def test_failed_detail_is_not_prepared_for_database(self) -> None:
        rows = [
            {"safetyCode": "FAILED", "_captureError": "detail failed"},
            {"safetyCode": "OK", "theme": "complete"},
        ]
        records = prepare_records(rows)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0][1], "OK")

    def test_excel_sheet_name_validation(self) -> None:
        self.assertEqual(validate_sheet_name("日常管理"), "日常管理")
        with self.assertRaises(ValueError):
            validate_sheet_name("bad/name")
        with self.assertRaises(ValueError):
            validate_sheet_name("x" * 32)

    def test_excel_file_can_be_reopened(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "result.xlsx"
            write_xlsx(output, '日常"管理', [["CODE-1", "单位"]])
            workbook = load_workbook(output, read_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ['日常"管理'])
                self.assertEqual(workbook.active["A2"].value, "CODE-1")
            finally:
                workbook.close()

    def test_spreadsheet_outputs_preserve_raw_values_and_neutralize_xlsx_formulas(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            output = root / "result.xlsx"
            write_xlsx(output, "日常管理", [["=1+1", "+SUM(A1:A2)"]])
            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(workbook.active["A2"].value, "=1+1")
                self.assertEqual(workbook.active["A2"].data_type, "s")
                self.assertEqual(workbook.active["B2"].value, "+SUM(A1:A2)")
                self.assertEqual(workbook.active["B2"].data_type, "s")
            finally:
                workbook.close()

            csv_path = root / "result.csv"
            safe_csv_path = root / "result-excel-safe.csv"
            rows = [{"theme": "@SUM(A1:A2)", "remark": "-现场整改"}]
            OutputWriter.write_csv(csv_path, rows)
            OutputWriter.write_csv(safe_csv_path, rows, spreadsheet_safe=True)
            with csv_path.open(encoding="utf-8-sig", newline="") as file:
                row = next(csv.DictReader(file))
                self.assertEqual(row["theme"], "@SUM(A1:A2)")
                self.assertEqual(row["remark"], "-现场整改")
            with safe_csv_path.open(encoding="utf-8-sig", newline="") as file:
                row = next(csv.DictReader(file))
                self.assertEqual(row["theme"], "'@SUM(A1:A2)")
                self.assertEqual(row["remark"], "'-现场整改")

    def test_detail_failure_does_not_overwrite_completed_batch(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            existing_details = root / "management-api-details.json"
            existing_details.write_text('{"rows":[{"safetyCode":"OLD"}]}', encoding="utf-8")
            result = CaptureResult(
                rows=[{"id": "1"}],
                details=[],
                failed_details=[{"id": "1", "_captureError": "timeout"}],
                request_body={},
                statistics=CaptureStatistics(
                    matched_rows=1,
                    details_requested=1,
                    details_failed=1,
                ),
            )
            args = Namespace(
                start="2026-07-01",
                end="2026-07-01",
                page_size=100,
                max_pages=10,
                request_timeout=1,
                retries=1,
                retry_backoff=0,
                max_detail_failures=0,
                out=root,
                headers_file=root / "headers.json",
                body_file=None,
                token_source=root,
                date_field="createTime",
                no_empty_page_stop=False,
                no_total_validation=False,
                allow_partial=False,
                details=True,
                verbose=False,
            )
            with patch(
                "data_acquisition.capture_daily_range.setup_logging",
                return_value=Mock(),
            ), patch(
                "data_acquisition.capture_daily_range.TokenProvider.resolve",
                return_value="token",
            ), patch(
                "data_acquisition.capture_daily_range.CaptureService.capture",
                return_value=result,
            ), patch(
                "data_acquisition.capture_daily_range.OutputWriter.write"
            ) as write:
                with self.assertRaisesRegex(RuntimeError, "failure threshold"):
                    run_capture(args, load_runtime_config())
            write.assert_not_called()
            self.assertEqual(
                existing_details.read_text(encoding="utf-8"),
                '{"rows":[{"safetyCode":"OLD"}]}',
            )
            attempts = list((root / "failed-attempts").iterdir())
            self.assertEqual(len(attempts), 1)
            self.assertTrue((attempts[0] / "management-api-detail-failures.json").is_file())

    def test_incomplete_list_is_diagnostic_unless_explicitly_allowed(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            result = CaptureResult(
                rows=[
                    {
                        "id": "1",
                        "safetyCode": "CODE-1",
                        "createTime": "2026-07-01 10:00:00",
                    }
                ],
                details=[],
                failed_details=[],
                request_body={},
                statistics=CaptureStatistics(
                    source_total_reported=2,
                    source_rows_fetched=1,
                    unique_rows=1,
                    matched_rows=1,
                    total_consistent=False,
                    incomplete_reasons=["API total mismatch"],
                ),
            )
            args = Namespace(
                start="2026-07-01",
                end="2026-07-01",
                page_size=100,
                max_pages=10,
                request_timeout=1,
                retries=1,
                retry_backoff=0,
                max_detail_failures=0,
                out=root / "rejected",
                headers_file=root / "headers.json",
                body_file=None,
                token_source=root,
                date_field="createTime",
                no_empty_page_stop=False,
                no_total_validation=False,
                allow_partial=False,
                details=False,
                verbose=False,
                browser_bridge_url="",
            )
            args.out.mkdir()
            existing = args.out / "management-api-all.json"
            existing.write_text("existing-complete-batch", encoding="utf-8")
            with patch(
                "data_acquisition.capture_daily_range.TokenProvider.resolve",
                return_value="token",
            ), patch(
                "data_acquisition.capture_daily_range.CaptureService.capture",
                return_value=result,
            ):
                with self.assertRaisesRegex(RuntimeError, "diagnostic files were retained"):
                    run_capture(args, load_runtime_config())
                self.assertEqual(existing.read_text(encoding="utf-8"), "existing-complete-batch")
                attempts = list((args.out / "failed-attempts").glob("capture-*"))
                self.assertEqual(len(attempts), 1)
                self.assertTrue((attempts[0] / "management-api-all.json").is_file())
                (args.out / "capture.log").replace(args.out / "capture-closed.log")

                args.out = root / "allowed"
                args.allow_partial = True
                paths = run_capture(args, load_runtime_config())

            self.assertEqual(paths["list_json"], args.out / "management-api-all.json")

    def test_detail_batch_includes_separate_failure_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            (root / "management-api-details.json").write_text(
                json.dumps({"rows": [{"safetyCode": "OK"}]}),
                encoding="utf-8",
            )
            (root / "management-api-detail-failures.json").write_text(
                json.dumps(
                    {"rows": [{"safetyCode": "FAILED", "_captureError": "timeout"}]}
                ),
                encoding="utf-8",
            )
            complete, failed = load_detail_batch(root)
            self.assertEqual([row["safetyCode"] for row in complete], ["OK"])
            self.assertEqual([row["safetyCode"] for row in failed], ["FAILED"])


if __name__ == "__main__":
    unittest.main()
